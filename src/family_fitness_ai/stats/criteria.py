"""등급 기준표(xlsx) → 문턱 표.

기준표는 연령대별 시트 하나씩이고, 등급·성별이 병합 셀로 묶여 있다. 여기서 하는 일은
그 격자를 한 행 = 한 문턱으로 펴는 것뿐이다. 방향(클수록 좋은가)은 기준표에 없으므로
`items.Item.lower_is_better` 가 정본이다 (docs/02 §5.3).
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

import openpyxl

from . import items as I

_GRADE = re.compile(r"([123])\s*등급")
_SEX = {"남": "M", "여": "F"}
_AGE_MAX = 120  # '85이상' 같은 열린 구간의 상한. 분포 산출에서만 쓰는 값이다.


@dataclass(frozen=True)
class BodyRange:
    """신체조성 기준. 문턱이 아니라 **구간**이라 별도 표로 둔다 (docs/dev/AI-2 §6).

    `lo` 이상 `hi` 미만이면 통과다. 한쪽만 있는 칸(`< 24.2`)은 다른 쪽이 `None` 이다.
    시트의 `초과`/`이상` 구분은 경계값 한 점의 차이라 구별하지 않는다.
    """

    age_group: str
    sex: str
    age_lo: int
    age_hi: int
    item_code: str
    lo: float | None
    hi: float | None

    def contains(self, value: float) -> bool:
        return (self.lo is None or value >= self.lo) and (self.hi is None or value < self.hi)


# 시트의 신체조성 열 머리글 → 항목 코드. 연령대마다 재는 것이 다르다.
BODY_COLUMNS = {"BMI": "018", "체지방률": "003", "WHtR": "042", "허리둘레": "042"}

_RANGE_BOTH = re.compile(r"([\d.]+)\s*%?\s*(?:이상|초과)\s*([\d.]+)\s*%?\s*미만")
_RANGE_UPPER = re.compile(r"^<\s*([\d.]+)$")


def parse_body_range(raw: object) -> tuple[float | None, float | None] | None:
    """'18.5이상 25미만' · '7%초과 27%미만' · '< 24.2' · '< .50' 을 구간으로."""
    if raw is None:
        return None
    text = " ".join(str(raw).split())
    if not text:
        return None
    if m := _RANGE_BOTH.search(text):
        return (float(m.group(1)), float(m.group(2)))
    if m := _RANGE_UPPER.match(text):
        return (None, float(m.group(1)))
    return None


@dataclass(frozen=True)
class Threshold:
    age_group: str
    sex: str
    age_lo: int
    age_hi: int
    item_code: str
    grade: int
    value: float


def parse_age_band(raw: object) -> tuple[int, int] | None:
    """'48~53' · 11 · '19~24' · '85이상' → (lo, hi). 유아기는 개월이다."""
    if raw is None:
        return None
    if isinstance(raw, (int, float)):
        n = int(raw)
        return (n, n)
    text = str(raw).strip().replace(" ", "")
    if not text:
        return None
    if "이상" in text:
        lo = int(re.sub(r"\D", "", text))
        return (lo, _AGE_MAX)
    m = re.match(r"^(\d+)[~\-](\d+)$", text)
    if m:
        return (int(m.group(1)), int(m.group(2)))
    if text.isdigit():
        n = int(text)
        return (n, n)
    return None


def _as_float(v: object) -> float | None:
    """숫자 칸만 받는다. '< 24.2' 같은 구간 형태는 점수화 대상이 아니다."""
    if isinstance(v, (int, float)):
        return float(v)
    return None


def load_xlsx(xlsx_path: str | Path) -> list[Threshold]:
    """기준표 원본(xlsx)에서 읽는다. 원본이 갱신될 때만 쓴다."""
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    out: list[Threshold] = []

    for sheet in wb.worksheets:
        age_group = I.SHEET_TO_AGE_GROUP.get(sheet.title)
        if age_group not in I.SCORED_AGE_GROUPS:
            continue
        columns = I.CRITERIA_COLUMNS[age_group]

        rows = list(sheet.iter_rows(values_only=True))
        header = [I.normalise_header(v) for v in rows[2]]
        col_of: dict[int, tuple[str, ...]] = {
            idx: codes for idx, text in enumerate(header) if (codes := columns.get(text))
        }
        missing = set(columns) - {header[i] for i in col_of}
        if missing:
            raise ValueError(f"{sheet.title}: 기준표 열을 찾지 못했다 — {sorted(missing)}")

        grade = sex = None
        for row in rows[3:]:
            if m := _GRADE.search(str(row[0] or "")):
                grade = int(m.group(1))
            if s := _SEX.get(str(row[1] or "").strip()):
                sex = s
            band = parse_age_band(row[2])
            if grade is None or sex is None or band is None:
                continue
            for idx, codes in col_of.items():
                value = _as_float(row[idx])
                if value is None:
                    continue
                for code in codes:
                    out.append(Threshold(age_group, sex, band[0], band[1], code, grade, value))
    return out


def load_body_ranges_xlsx(xlsx_path: str | Path) -> list[BodyRange]:
    """기준표 원본에서 신체조성 구간을 읽는다. **3등급 행에만 있다.**

    시트에서 1·2등급 행의 BMI·체지방률 칸은 비어 있다 — 신체조성은 3등급 판정에만
    쓰인다 (docs/dev/AI-2 §6). 병합 셀이라 값을 각 행에 채워 넣고 읽는다.
    """
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    out: list[BodyRange] = []

    for sheet in wb.worksheets:
        # 열 매핑이 있는 시트만 읽는다. `SCORED_AGE_GROUPS` 를 쓰지 않는 것은 그 상수의
        # 자리가 갈래마다 다르기 때문이다 — 여기서 필요한 것은 "읽을 열을 아는 시트"다.
        age_group = I.SHEET_TO_AGE_GROUP.get(sheet.title)
        if age_group is None or age_group not in I.CRITERIA_COLUMNS:
            continue
        grid = [[c.value for c in row] for row in sheet.iter_rows()]
        for rng in sheet.merged_cells.ranges:
            value = grid[rng.min_row - 1][rng.min_col - 1]
            for r in range(rng.min_row - 1, rng.max_row):
                for c in range(rng.min_col - 1, rng.max_col):
                    grid[r][c] = value

        header = [I.normalise_header(v) for v in grid[2]]
        col_of = {
            idx: code
            for idx, text in enumerate(header)
            for key, code in BODY_COLUMNS.items()
            if key in text
        }
        for row in grid[3:]:
            if not (row[0] and "3등급" in str(row[0])):
                continue
            sex = _SEX.get(str(row[1] or "").strip())
            band = parse_age_band(row[2])
            if sex is None or band is None:
                continue
            for idx, code in col_of.items():
                parsed = parse_body_range(row[idx])
                if parsed is None:
                    continue
                out.append(BodyRange(age_group, sex, band[0], band[1], code, *parsed))
    return out


def body_ranges_to_frame(ranges: list[BodyRange]):
    import pandas as pd

    return pd.DataFrame(
        [
            {
                "age_group": r.age_group,
                "sex": r.sex,
                "age_lo": r.age_lo,
                "age_hi": r.age_hi,
                "age_unit": "개월" if r.age_group == "유아기" else "세",
                "item_code": r.item_code,
                "item_name": {"018": "BMI", "003": "체지방률", "042": "WHtR"}[r.item_code],
                "range_lo": "" if r.lo is None else r.lo,
                "range_hi": "" if r.hi is None else r.hi,
            }
            for r in ranges
        ]
    ).sort_values(["age_group", "sex", "age_lo", "item_code"])


def load_body_ranges_csv(csv_path: str | Path) -> list[BodyRange]:
    import csv as _csv

    def num(text: str) -> float | None:
        return float(text) if text not in ("", None) else None

    with open(csv_path, encoding="utf-8-sig", newline="") as fh:
        return [
            BodyRange(
                age_group=r["age_group"],
                sex=r["sex"],
                age_lo=int(r["age_lo"]),
                age_hi=int(r["age_hi"]),
                item_code=r["item_code"],
                lo=num(r["range_lo"]),
                hi=num(r["range_hi"]),
            )
            for r in _csv.DictReader(fh)
        ]


def to_frame(thresholds: list[Threshold]):
    """문턱을 평평한 표로. 기준표 원본(xlsx) 없이도 값을 확인할 수 있게 한다."""
    import pandas as pd

    return pd.DataFrame(
        [
            {
                "age_group": t.age_group,
                "sex": t.sex,
                "age_lo": t.age_lo,
                "age_hi": t.age_hi,
                "age_unit": "개월" if t.age_group == "유아기" else "세",
                "item_code": t.item_code,
                "item_name": I.ITEMS[t.item_code].name,
                "unit": I.ITEMS[t.item_code].unit,
                "lower_is_better": int(I.ITEMS[t.item_code].lower_is_better),
                "grade": t.grade,
                "threshold": t.value,
            }
            for t in thresholds
        ]
    ).sort_values(["age_group", "sex", "age_lo", "item_code", "grade"])


def load_csv(csv_path: str | Path) -> list[Threshold]:
    """레포에 커밋된 문턱 표에서 읽는다. 원본 xlsx 가 없어도 산출이 돈다."""
    import csv as _csv

    with open(csv_path, encoding="utf-8-sig", newline="") as fh:
        return [
            Threshold(
                age_group=r["age_group"],
                sex=r["sex"],
                age_lo=int(r["age_lo"]),
                age_hi=int(r["age_hi"]),
                item_code=r["item_code"],
                grade=int(r["grade"]),
                value=float(r["threshold"]),
            )
            for r in _csv.DictReader(fh)
        ]


def load(path: str | Path) -> list[Threshold]:
    """확장자로 원본(xlsx)과 커밋본(csv)을 가른다."""
    suffix = Path(path).suffix.lower()
    if suffix == ".csv":
        return load_csv(path)
    if suffix in (".xlsx", ".xlsm"):
        return load_xlsx(path)
    raise ValueError(f"기준표는 .xlsx 또는 .csv 여야 한다: {path}")
